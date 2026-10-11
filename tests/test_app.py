import io
import json
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import create_app
from conftest import FakeFetcher
from test_sync import base_pages


def client_for(env, fixtures, **kw):
    pages = base_pages(fixtures)
    fetcher = FakeFetcher(pages)
    app = create_app(fetcher_factory=lambda: fetcher, **kw)
    c = TestClient(app)
    return c, app, fetcher


def wait_idle(app):
    for _ in range(100):
        if not app.state.syncer.running:
            return
        time.sleep(0.05)


def test_dashboard_empty_then_full(env, fixtures):
    c, app, _ = client_for(env, fixtures)
    r = c.get("/")
    assert r.status_code == 200
    assert "No results yet" in r.text
    app.state.syncer.sync("manual")
    r = c.get("/")
    assert r.status_code == 200
    assert "Latest weekend" in r.text
    assert "Ava" in r.text and "Max" in r.text and "Sam" in r.text
    assert "New PB" in r.text
    assert "Half marathon (11)" in r.text
    assert "25" in r.text and "of 46" in r.text      # gender position
    assert "6:22 /km" in r.text and "Average pace" in r.text
    assert c.get("/healthz").json()["ok"]


def test_refresh_and_cooldown(env, fixtures, monkeypatch):
    monkeypatch.setenv("REFRESH_COOLDOWN", "300")
    c, app, fetcher = client_for(env, fixtures)
    r = c.post("/api/refresh")
    assert r.status_code == 202
    wait_idle(app)
    s = c.get("/api/status").json()
    assert s["last_attempt"]["new_results"] == 8
    assert s["cooldown_seconds"] > 0
    r = c.post("/api/refresh")
    assert r.status_code == 429
    assert "Try again" in r.json()["message"]


def test_exports(env, fixtures):
    c, app, _ = client_for(env, fixtures)
    app.state.syncer.sync("manual")
    r = c.get("/export/results.csv")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    lines = r.content.decode("utf-8-sig").strip().splitlines()
    assert lines[0].startswith("Runner,parkrun ID,Date")
    assert len(lines) == 9
    j = c.get("/export/results.json?runner=Ava").json()
    assert len(j["results"]) == 3 and j["results"][0]["gender_position"] == 25
    assert j["results"][0]["pace"] == "6:22" and j["results"][0]["pace_seconds_per_km"] == 382.5
    assert "Pace (min/km)" in lines[0]
    x = c.get("/export/results.xlsx?distance=2k")
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.active.max_row == 1 + 7
    assert c.get("/export/results.pdf").status_code == 404


def test_api_for_home_assistant(env, fixtures):
    c, app, _ = client_for(env, fixtures)
    app.state.syncer.sync("manual")
    latest = c.get("/api/latest").json()["latest"]
    assert latest["Ava"]["time"] == "12:45" and latest["Sam"] is None
    summary = c.get("/api/summary").json()["runners"]
    assert summary["Ava"]["distances"]["2k"]["runs_to_next_milestone"] == 8
    assert summary["Max"]["distances"]["5k"]["pb"] == "29:58"
    assert summary["Max"]["distances"]["5k"]["pb_pace"] == "6:00 /km"   # 29:58 / 5 km


def test_import_endpoint(env, fixtures):
    c, app, _ = client_for(env, fixtures)
    r = c.post("/api/import", files={"file": ("p.html", fixtures["profile_junior.html"], "text/html")})
    assert r.status_code == 200 and r.json()["new_results"] == 3
    r = c.post("/api/import", files={"file": ("x.html", "<html>nope</html>", "text/html")})
    assert r.status_code == 422


def test_removed_runner_hidden(env, fixtures):
    c, app, _ = client_for(env, fixtures)
    app.state.syncer.sync("manual")
    # drop Max from the config and restart
    (env / "runners.yaml").write_text("runners:\n  - name: Ava\n    id: A1000001\n")
    c2 = TestClient(create_app(fetcher_factory=lambda: FakeFetcher({})))
    r = c2.get("/")
    assert "Max" not in r.text and "Ava" in r.text
