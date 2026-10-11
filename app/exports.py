"""Raw data exports: CSV, JSON and Excel."""
from __future__ import annotations

import csv
import io
import json

from .stats import Result

COLUMNS = [
    ("runner", "Runner"), ("parkrun_id", "parkrun ID"), ("date", "Date"), ("event", "Event"),
    ("run_number", "Run number"), ("distance", "Distance"), ("time", "Time"),
    ("time_seconds", "Time (seconds)"), ("position", "Position"), ("finishers", "Finishers"),
    ("gender", "Gender"), ("gender_position", "Gender position"), ("gender_count", "Gender finishers"),
    ("age_group", "Age group"), ("age_grade", "Age grade %"), ("new_pb", "New PB"),
    ("first_at_distance", "First at distance"), ("event_pb", "Event PB"), ("milestone", "Milestone"),
    ("results_url", "Results page"),
]


def filtered(results: list[Result], runner: str | None, distance: str | None) -> list[dict]:
    rows = [r.as_dict() for r in sorted(results, key=lambda r: (r.run_date, r.runner), reverse=True)]
    if runner:
        rows = [r for r in rows if r["runner"].lower() == runner.lower()]
    if distance:
        rows = [r for r in rows if r["distance"] == distance]
    return rows


def to_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([label for _, label in COLUMNS])
    for r in rows:
        w.writerow(["" if r[k] is None else r[k] for k, _ in COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")   # BOM so Excel reads UTF-8 correctly


def to_json(rows: list[dict]) -> bytes:
    return json.dumps({"results": rows}, indent=2).encode("utf-8")


def to_xlsx(rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append([label for _, label in COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r[k] for k, _ in COLUMNS])
    date_col = [k for k, _ in COLUMNS].index("date") + 1
    for row in ws.iter_rows(min_row=2, min_col=date_col, max_col=date_col):
        for c in row:
            from datetime import date
            c.value = date.fromisoformat(c.value) if c.value else None
            c.number_format = "dd/mm/yyyy"
    for i, (k, label) in enumerate(COLUMNS, start=1):
        width = max([len(label)] + [len(str(r[k])) for r in rows if r[k] is not None][:200])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
